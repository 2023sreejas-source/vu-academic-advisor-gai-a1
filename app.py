# ============================================================
# VU AI ACADEMIC ADVISOR — Vidyashilp University
# DATA308 Generative AI Assignment 2
# ============================================================

import streamlit as st
from groq import Groq
from pypdf import PdfReader

import os
import re
import math
import json
import requests
import html

from collections import Counter


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Academic Advisor — Vidyashilp University",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ============================================================
# CSS
# ============================================================

st.markdown(
    """
    <style>

    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

    html, body, [class*="css"] {
        font-family: 'Google Sans', 'Google Sans Text',
        'Product Sans', 'Inter', Arial, sans-serif;
    }

    .stApp,
    [data-testid="stAppViewContainer"],
    [data-testid="stAppViewBlockContainer"],
    [data-testid="block-container"] {
        background: #ffffff !important;
    }

    [data-testid="stMainBlockContainer"] h1,
    [data-testid="stMainBlockContainer"] h2,
    [data-testid="stMainBlockContainer"] h3,
    [data-testid="stMainBlockContainer"] p,
    [data-testid="stMainBlockContainer"] strong {
        color: #0b4f8a !important;
    }

    /* Sidebar */

    [data-testid="stSidebar"] {
        background: #082f5b !important;
        border-right: 1px solid #dbe4ee !important;
    }

    [data-testid="stSidebar"] > div:first-child {
        background: #082f5b !important;
    }

    [data-testid="stSidebar"] * {
        color: #ffffff !important;
    }

    [data-testid="stSidebar"] hr {
        border-color: rgba(255,255,255,0.20) !important;
    }

    [data-testid="stSidebar"] label {
        color: #ffffff !important;
        font-weight: 500 !important;
    }

    [data-testid="stSidebar"] input {
        color: #172033 !important;
        background: #ffffff !important;
    }

    [data-testid="stSidebar"] [data-baseweb="select"] > div,
    [data-testid="stSidebar"] [data-baseweb="input"] > div {
        background: #ffffff !important;
        border-color: #d9e2ec !important;
        color: #172033 !important;
    }

    [data-testid="stSidebar"] [data-baseweb="select"] span,
    [data-testid="stSidebar"] [data-baseweb="input"] input {
        color: #172033 !important;
    }

    [data-testid="stSidebar"] div[data-testid="stButton"] > button {
        background: transparent !important;
        border: 1px solid rgba(255,255,255,0.70) !important;
        color: #ffffff !important;
    }

    [data-testid="stSidebar"] div[data-testid="stButton"] > button:hover {
        background: rgba(255,255,255,0.10) !important;
        border-color: #ffffff !important;
    }

    /* Chat input */

    [data-testid="stBottom"],
    [data-testid="stBottom"] > div {
        background: #ffffff !important;
        border-top: 1px solid #e7ebf0 !important;
        padding: 10px 16px !important;
    }

    [data-testid="stChatInputContainer"],
    [data-testid="stChatInputContainer"] > div {
        background: #ffffff !important;
        border: 1.5px solid #d9e1ea !important;
        border-radius: 26px !important;
        box-shadow: 0 2px 8px rgba(8,47,91,0.07) !important;
    }

    [data-testid="stChatInputContainer"] textarea {
        color: #172033 !important;
        background: transparent !important;
        font-size: 15px !important;
    }

    /* Buttons */

    div[data-testid="stButton"] > button {
        background: #ffffff !important;
        border: 1px solid #d9e2ec !important;
        border-radius: 20px !important;
        color: #0b4f8a !important;
        font-size: 14px !important;
        font-weight: 500 !important;
    }

    div[data-testid="stButton"] > button:hover {
        border-color: #0b4f8a !important;
        background: #f3f7fb !important;
    }

    /* Welcome */

    .info-box {
        background: #ffffff;
        padding: 17px 20px;
        border-radius: 14px;
        border: 1px solid #e1e7ee;
        border-left: 4px solid #0b4f8a;
        margin-bottom: 18px;
        box-shadow: 0 2px 8px rgba(8,47,91,0.05);
    }

    .info-box h3 {
        color: #172033 !important;
    }

    .info-box p {
        color: #334155 !important;
        font-size: 15px !important;
        line-height: 1.65 !important;
    }

    /* Chat */

    .bubble-user {
        display: flex;
        justify-content: flex-end;
        margin: 11px 0 14px;
    }

    .bubble-user-inner {
        background: #0b4f8a;
        color: #ffffff !important;
        border-radius: 18px 18px 5px 18px;
        padding: 11px 17px;
        max-width: 68%;
        font-size: 15px;
        line-height: 1.65;
        word-wrap: break-word;
    }

    .bubble-bot {
        display: flex;
        justify-content: flex-start;
        margin: 11px 0 14px;
    }

    .bubble-bot-inner {
        background: #f7f9fc;
        color: #172033;
        border: 1px solid #dfe6ee;
        border-radius: 18px 18px 18px 5px;
        padding: 12px 17px;
        max-width: 76%;
        font-size: 15px;
        line-height: 1.7;
        word-wrap: break-word;
    }

    /* Profile */

    .profile-card {
        background: rgba(255,255,255,0.10);
        border: 1px solid rgba(255,255,255,0.20);
        border-radius: 10px;
        padding: 10px 12px;
        margin: 8px 0 14px;
    }

    .profile-card-title {
        color: #ffffff !important;
        font-weight: 700;
        font-size: 14px;
    }

    .profile-card-text {
        color: rgba(255,255,255,0.85) !important;
        font-size: 12px;
        line-height: 1.5;
    }

    .footer-text {
        text-align: center;
        color: #a1aab5;
        font-size: 11px;
        padding-top: 20px;
    }

    #MainMenu,
    footer {
        visibility: hidden;
    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# CONSTANTS
# ============================================================

MODEL = "llama-3.3-70b-versatile"

ACADEMIC_EXTENSIONS = (
    ".pdf",
    ".txt",
    ".csv",
    ".xlsx",
    ".xls"
)

STUDENT_DATABASE_FILE = "synthetic_students.json"

EXCLUDED_FILES = (
    "advisor_eval",
    "eval_results",
    "phase4",
    "summary_metrics",
    "website_sources"
)

DEFAULT_WEBSITE_SOURCES = [
    {
        "name": "VU Official Website",
        "url": "https://vidyashilp.edu.in/"
    },
    {
        "name": "VU Admissions",
        "url": "https://vidyashilp.edu.in/admission-enquiryform/"
    },
    {
        "name": "VU Contact",
        "url": "https://vidyashilp.edu.in/contact/"
    },
    {
        "name": "VU B.Tech AI/ML",
        "url": "https://vidyashilp.edu.in/schools/b_tech_ai_ml/"
    },
    {
        "name": "VU BMS",
        "url": "https://vidyashilp.edu.in/schools/b-m-s-hons-hons-with-research-in-digital-business/"
    },
    {
        "name": "VU B.A., LL.B. (Hons.)",
        "url": "https://vidyashilp.edu.in/ba-llb/"
    },
    {
        "name": "VU B.M.S., LL.B. (Hons.)",
        "url": "https://vidyashilp.edu.in/schools/bachelor-of-management-studies-bachelor-of-laws-hons/"
    }
]


# ============================================================
# GREETINGS / SOCIAL
# ============================================================

SOCIAL_RE = re.compile(
    r"^\s*(h+i+|h+e+l+o+|hey+|hola|yo+|sup|bro|wsp|wsup|"
    r"what'?s\s*up|wassup|good\s*(morning|afternoon|evening|night)|"
    r"namaste|namaskaram|thank(s|\s*you|u)|thx|ty|ok(ay)?|"
    r"got\s*it|sure|great|nice|cool|bye|goodbye|see\s*you|"
    r"take\s*care|cya)\s*[!.?]*\s*$",
    re.I
)

GREETING_REPLY = (
    "Hey! 👋 I'm VU's AI Academic Advisor. "
    "I can help with attendance, credits, prerequisites, registration, "
    "course eligibility, progression rules and more. "
    "What would you like to know?"
)

THANKS_REPLY = (
    "Happy to help! 😊 Feel free to ask anything else "
    "about your academics at VU."
)

OUT_OF_SCOPE_RE = re.compile(
    r"\b(weather|temperature|rain|forecast|cricket|football|soccer|"
    r"match|movie|movies|song|music|restaurant|recipe|stock market|"
    r"politics|election|celebrity|who will win|score today|bitcoin|"
    r"share price)\b",
    re.I
)

PERSONAL_RE = re.compile(
    r"\b(your mom|your mother|your dad|your father|your family|"
    r"your girlfriend|your boyfriend|are you dumb|are u dumb|"
    r"are you stupid|are u stupid)\b",
    re.I
)

OUT_OF_SCOPE_REPLY = (
    "I’m VU’s AI Academic Advisor, so I can help with Vidyashilp "
    "University academic and programme-related questions, but not that topic. 😊"
)

PERSONAL_REPLY = (
    "I’m an AI academic advisor, so I don’t have a personal family "
    "or personal life. But I can definitely help with your VU academic questions. 😊"
)


# ============================================================
# SESSION STATE
# ============================================================

if "messages" not in st.session_state:
    st.session_state.messages = []

if "academic_kb" not in st.session_state:
    st.session_state.academic_kb = []

if "website_kb" not in st.session_state:
    st.session_state.website_kb = []

if "knowledge_loaded" not in st.session_state:
    st.session_state.knowledge_loaded = False

if "website_loaded" not in st.session_state:
    st.session_state.website_loaded = False

if "selected_student_id" not in st.session_state:
    st.session_state.selected_student_id = "None"


# ============================================================
# SYNTHETIC STUDENT DATABASE
# ============================================================

SENSITIVE_KEYS = {
    "email",
    "personal_email",
    "phone",
    "mobile",
    "address",
    "home_address",
    "date_of_birth",
    "dob",
    "password",
    "contact_number"
}


def sanitize_student(student):
    """
    Remove obvious personal/contact fields from synthetic profiles.
    """

    if not isinstance(student, dict):
        return {}

    cleaned = {}

    for key, value in student.items():
        if str(key).lower() not in SENSITIVE_KEYS:
            cleaned[key] = value

    return cleaned


def load_student_database():
    """
    Load synthetic student profiles.

    Supports:
    - a JSON list
    - {"students": [...]}
    - {"profiles": [...]}
    - {"student_profiles": [...]}
    - {"data": [...]}

    The database is only student context.
    It is NOT an authority for university rules.
    """

    if not os.path.exists(STUDENT_DATABASE_FILE):
        return {}

    try:
        with open(
            STUDENT_DATABASE_FILE,
            "r",
            encoding="utf-8"
        ) as f:
            data = json.load(f)

    except Exception:
        return {}

    if isinstance(data, list):
        students = data

    elif isinstance(data, dict):

        students = (
            data.get("students")
            or data.get("profiles")
            or data.get("student_profiles")
            or data.get("data")
            or []
        )

    else:
        return {}

    database = {}

    for student in students:

        if not isinstance(student, dict):
            continue

        student = sanitize_student(student)

        student_id = str(
            student.get("student_id", "")
        ).strip()

        if student_id:
            database[student_id] = student

    return database


STUDENT_DATABASE = load_student_database()


# ============================================================
# HEADER & LOGO
# ============================================================

def find_logo():

    for filename in os.listdir("."):

        if filename.startswith("."):
            continue

        lower = filename.lower()

        if (
            lower.endswith((".png", ".jpg", ".jpeg"))
            and "logo" in lower
        ):
            return filename

    for filename in os.listdir("."):

        lower = filename.lower()

        if lower.endswith((".png", ".jpg", ".jpeg")):
            return filename

    return None


logo_path = find_logo()

header_col1, header_col2 = st.columns(
    [1, 7],
    vertical_alignment="center"
)

with header_col1:

    if logo_path:
        st.image(
            logo_path,
            width=90
        )

with header_col2:

    st.markdown(
        """
        <div style="padding: 5px 0;">
            <span style="
                font-size: 28px;
                font-weight: 700;
                color: #0b4f8a !important;
                display: block;
                line-height: 1.2;">
                Vidyashilp University
            </span>

            <span style="
                font-size: 15px;
                font-weight: 600;
                color: #475569 !important;
                display: block;
                margin-top: 4px;">
                AI Academic Advisor · Online
            </span>
        </div>
        """,
        unsafe_allow_html=True
    )

st.divider()


# ============================================================
# SIDEBAR — STUDENT PROFILE
# ============================================================

with st.sidebar:

    st.markdown("### 🎓 Student Profile")

    st.markdown(
        "Select a synthetic student for A2 testing, "
        "or use the manual profile."
    )

    st.markdown("---")

    # --------------------------------------------------------
    # Synthetic student search
    # --------------------------------------------------------

    student_search = st.text_input(
        "Search Student ID",
        placeholder="e.g. STU001"
    )

    all_student_ids = list(
        STUDENT_DATABASE.keys()
    )

    if student_search.strip():

        search_value = student_search.strip().lower()

        filtered_student_ids = [
            sid
            for sid in all_student_ids
            if search_value in sid.lower()
        ]

    else:

        filtered_student_ids = all_student_ids

    student_options = ["None"] + filtered_student_ids

    if (
        st.session_state.selected_student_id
        not in student_options
    ):
        st.session_state.selected_student_id = "None"

    selected_student_id = st.selectbox(
        "Synthetic Student",
        student_options,
        index=student_options.index(
            st.session_state.selected_student_id
        ),
        help="Synthetic profiles created specifically for A2 testing."
    )

    st.session_state.selected_student_id = selected_student_id

    selected_student = None

    if selected_student_id != "None":

        selected_student = STUDENT_DATABASE.get(
            selected_student_id
        )

    # --------------------------------------------------------
    # Selected synthetic profile
    # --------------------------------------------------------

    if selected_student:

        student_program = selected_student.get(
            "program",
            selected_student.get(
                "programme",
                "Not specified"
            )
        )

        student_semester = selected_student.get(
            "semester",
            "Not specified"
        )

        scenario = (
            selected_student.get("scenario")
            or selected_student.get("scenario_description")
            or selected_student.get("testing_purpose")
            or selected_student.get("student_scenario")
            or selected_student.get("scenario_type")
            or "Synthetic A2 test scenario"
        )

        if isinstance(scenario, list):

            scenario_text = ", ".join(
                str(x)
                for x in scenario[:3]
            )

        else:

            scenario_text = str(scenario)

        st.markdown(
            f"""
            <div class="profile-card">

                <div class="profile-card-title">
                    Active Synthetic Profile
                </div>

                <div class="profile-card-text">

                    ID:
                    {html.escape(str(selected_student_id))}
                    <br>

                    Programme:
                    {html.escape(str(student_program))}
                    <br>

                    Semester:
                    {html.escape(str(student_semester))}
                    <br>

                    Scenario:
                    {html.escape(scenario_text[:220])}

                </div>

            </div>
            """,
            unsafe_allow_html=True
        )

    elif not STUDENT_DATABASE:

        st.warning(
            "synthetic_students.json was not found "
            "or could not be read."
        )

    elif student_search and not filtered_student_ids:

        st.info(
            "No student ID matched your search."
        )

    st.markdown("---")

    # --------------------------------------------------------
    # Manual profile
    # --------------------------------------------------------

    st.markdown("**Manual Profile**")

    program = st.selectbox(
        "Program",
        [
            "Not specified",
            "B.Tech",
            "BMS",
            "BA LLB",
            "BMS LLB",
            "B.A. Economics",
            "B.A. Psychology",
            "B.Des"
        ],
        index=0
    )

    semester = st.selectbox(
        "Current Semester",
        [
            "Not specified",
            "1st Semester",
            "2nd Semester",
            "3rd Semester",
            "4th Semester",
            "5th Semester",
            "6th Semester",
            "7th Semester",
            "8th Semester",
            "9th Semester",
            "10th Semester"
        ],
        index=0
    )

    completed_credits = st.number_input(
        "Completed Credits",
        min_value=0,
        max_value=400,
        value=None,
        placeholder="Optional"
    )

    cgpa = st.number_input(
        "CGPA",
        min_value=0.0,
        max_value=10.0,
        value=None,
        step=0.1,
        placeholder="Optional"
    )

    st.markdown("---")

    if st.button(
        "🗑️ Clear Chat",
        use_container_width=True
    ):

        st.session_state.messages = []

        st.rerun()


# ============================================================
# API KEY
# ============================================================

def get_api_key():

    try:

        secret_key = st.secrets.get(
            "GROQ_API_KEY",
            ""
        )

        if secret_key:
            return secret_key

    except Exception:
        pass

    return os.getenv(
        "GROQ_API_KEY",
        ""
    )


api_key = get_api_key()


# ============================================================
# MANUAL PROFILE
# ============================================================

manual_student_profile = {
    "program": program,
    "semester": semester,
    "completed_credits": completed_credits,
    "cgpa": cgpa
}


# ============================================================
# ACTIVE PROFILE
# ============================================================

def get_active_student_profile():

    if selected_student:
        return selected_student

    return manual_student_profile


student_profile = get_active_student_profile()


# ============================================================
# PROFILE HELPERS
# ============================================================

def format_profile_value(value):

    if value is None:
        return "Not specified"

    if isinstance(value, list):

        if not value:
            return "None"

        return ", ".join(
            str(item)
            for item in value
        )

    if isinstance(value, dict):

        parts = []

        for key, item in value.items():

            parts.append(
                f"{key}: {item}"
            )

        return "; ".join(parts)

    return str(value)


def profile_to_text(profile):

    if not profile:
        return "No student profile available."

    important_fields = [
        "student_id",
        "program",
        "programme",
        "semester",
        "completed_credits",
        "credits",
        "cgpa",
        "minor",
        "passed_courses",
        "completed_courses",
        "failed_courses",
        "current_courses",
        "pending_courses",
        "interests",
        "academic_constraints",
        "known_prerequisites_status",
        "data_completeness",
        "scenario",
        "scenario_description",
        "student_scenario",
        "testing_purpose"
    ]

    lines = []

    for field in important_fields:

        if field not in profile:
            continue

        value = profile.get(field)

        if value is None:
            continue

        lines.append(
            f"- {field}: {format_profile_value(value)}"
        )

    return "\n".join(lines)


def profile_retrieval_terms(profile):

    if not profile:
        return ""

    terms = []

    fields = [
        "program",
        "programme",
        "minor",
        "passed_courses",
        "completed_courses",
        "failed_courses",
        "current_courses",
        "pending_courses"
    ]

    for field in fields:

        value = profile.get(field)

        if isinstance(value, list):

            terms.extend(
                str(item)
                for item in value
            )

        elif isinstance(value, str):

            terms.append(value)

    return " ".join(terms)


# ============================================================
# TEXT EXTRACTION
# ============================================================

def clean_text(text):

    if not text:
        return ""

    text = text.replace(
        "\x00",
        " "
    )

    text = re.sub(
        r"\s+",
        " ",
        text
    )

    return text.strip()


def guess_section(raw_text):

    for line in raw_text.splitlines():

        line = re.sub(
            r"\s+",
            " ",
            line
        ).strip()

        if not line or len(line) > 100:
            continue

        if re.match(
            r"^(?:\d+(?:\.\d+)*\s+.+|[A-Z][A-Z0-9 &:/()\-]{4,})$",
            line
        ):
            return line

        words = line.split()

        if (
            2 <= len(words) <= 10
            and not re.search(
                r"[.!?]$",
                line
            )
        ):

            title_case = sum(
                1
                for word in words
                if word[:1].isupper()
            )

            if title_case >= max(
                2,
                len(words) // 2
            ):
                return line

    return ""


def extract_pdf_pages(filepath):

    pages = []

    try:

        reader = PdfReader(filepath)

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            try:

                raw = page.extract_text() or ""

                text = clean_text(raw)

                if text:

                    pages.append(
                        {
                            "text": text,
                            "page": page_number,
                            "section": guess_section(raw)
                        }
                    )

            except Exception:
                continue

    except Exception:

        return []

    return pages


def extract_txt_text(filepath):

    try:

        with open(
            filepath,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            return clean_text(
                f.read()
            )

    except Exception:

        return ""


def extract_csv_text(filepath):

    try:

        with open(
            filepath,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            return clean_text(
                f.read()
            )

    except Exception:

        return ""


def extract_excel_sheets(filepath):

    try:

        from openpyxl import load_workbook

        wb = load_workbook(
            filepath,
            read_only=True,
            data_only=True
        )

        sheets = []

        for sheet in wb.worksheets:

            rows = []

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = [
                    str(v)
                    for v in row
                    if v is not None
                ]

                if values:

                    rows.append(
                        " | ".join(values)
                    )

            text = clean_text(
                "\n".join(rows)
            )

            if text:

                sheets.append(
                    {
                        "sheet": sheet.title,
                        "text": text
                    }
                )

        return sheets

    except Exception:

        return []


# ============================================================
# CHUNKING
# ============================================================

def chunk_text(
    text,
    chunk_size=150,
    overlap=30
):

    words = text.split()

    if not words:
        return []

    chunks = []

    start = 0

    while start < len(words):

        end = min(
            start + chunk_size,
            len(words)
        )

        chunk = " ".join(
            words[start:end]
        )

        if chunk.strip():

            chunks.append(
                chunk.strip()
            )

        if end >= len(words):
            break

        start = end - overlap

    return chunks


# ============================================================
# LOAD ACADEMIC DOCUMENTS
# ============================================================

def load_academic_documents():

    documents = []

    for filename in sorted(
        os.listdir(".")
    ):

        if filename.startswith("."):
            continue

        lower = filename.lower()

        if not lower.endswith(
            ACADEMIC_EXTENSIONS
        ):
            continue

        if any(
            excluded in lower
            for excluded in EXCLUDED_FILES
        ):
            continue

        filepath = os.path.join(
            ".",
            filename
        )

        ext = os.path.splitext(
            filename
        )[1].lower()

        # PDF

        if ext == ".pdf":

            for page_info in extract_pdf_pages(
                filepath
            ):

                chunks = chunk_text(
                    page_info["text"]
                )

                for i, chunk in enumerate(
                    chunks,
                    start=1
                ):

                    documents.append(
                        {
                            "text": chunk,
                            "source": filename,
                            "type": "academic",
                            "chunk": i,
                            "page": page_info.get("page"),
                            "section": page_info.get(
                                "section",
                                ""
                            )
                        }
                    )

        # Excel

        elif ext in (
            ".xlsx",
            ".xls"
        ):

            for sheet_info in extract_excel_sheets(
                filepath
            ):

                chunks = chunk_text(
                    sheet_info["text"]
                )

                for i, chunk in enumerate(
                    chunks,
                    start=1
                ):

                    documents.append(
                        {
                            "text": chunk,
                            "source": filename,
                            "type": "academic",
                            "chunk": i,
                            "sheet": sheet_info.get(
                                "sheet",
                                ""
                            )
                        }
                    )

        # TXT / CSV

        else:

            if ext == ".txt":
                text = extract_txt_text(filepath)
            else:
                text = extract_csv_text(filepath)

            for i, chunk in enumerate(
                chunk_text(text),
                start=1
            ):

                documents.append(
                    {
                        "text": chunk,
                        "source": filename,
                        "type": "academic",
                        "chunk": i
                    }
                )

    return documents


# ============================================================
# WEBSITE SOURCES
# ============================================================

def load_website_sources():

    source_file = "website_sources.json"

    if os.path.exists(
        source_file
    ):

        try:

            with open(
                source_file,
                "r",
                encoding="utf-8"
            ) as f:

                data = json.load(f)

            if (
                isinstance(data, list)
                and len(data) > 0
            ):

                return data

        except Exception:
            pass

    return DEFAULT_WEBSITE_SOURCES


@st.cache_data(
    ttl=3600,
    show_spinner=False
)
def fetch_webpage(url):

    try:

        headers = {
            "User-Agent":
            "Mozilla/5.0"
        }

        response = requests.get(
            url,
            headers=headers,
            timeout=20
        )

        response.raise_for_status()

        from bs4 import BeautifulSoup

        soup = BeautifulSoup(
            response.text,
            "html.parser"
        )

        for tag in soup(
            [
                "script",
                "style",
                "noscript",
                "svg",
                "iframe",
                "nav",
                "footer"
            ]
        ):

            tag.decompose()

        title = ""

        if soup.title:

            title = clean_text(
                soup.title.get_text()
            )

        sections = []

        current_section = (
            title
            or "VU Official Website"
        )

        for element in soup.find_all(
            [
                "h1",
                "h2",
                "h3",
                "p",
                "li"
            ]
        ):

            if element.name in (
                "h1",
                "h2",
                "h3"
            ):

                heading = clean_text(
                    element.get_text(" ")
                )

                if heading:
                    current_section = heading

            else:

                text = clean_text(
                    element.get_text(" ")
                )

                if text:

                    sections.append(
                        {
                            "section": current_section,
                            "text": text
                        }
                    )

        if sections:

            return {
                "title": title,
                "sections": sections
            }

        return {
            "title": title,
            "sections": [
                {
                    "section":
                        title or "VU Official Website",
                    "text":
                        clean_text(
                            soup.get_text(
                                separator=" "
                            )
                        )
                }
            ]
        }

    except Exception:

        return {
            "title": "",
            "sections": []
        }


def load_website_documents():

    documents = []

    for source in load_website_sources():

        name = source.get(
            "name",
            "VU Official Website"
        )

        url = source.get(
            "url",
            ""
        )

        if not url:
            continue

        page = fetch_webpage(
            url
        )

        for section_info in page.get(
            "sections",
            []
        ):

            text = section_info.get(
                "text",
                ""
            )

            if not text:
                continue

            for i, chunk in enumerate(
                chunk_text(
                    text,
                    chunk_size=180,
                    overlap=40
                ),
                start=1
            ):

                documents.append(
                    {
                        "text": chunk,
                        "source": name,
                        "url": url,
                        "title": page.get(
                            "title",
                            ""
                        ),
                        "section": section_info.get(
                            "section",
                            ""
                        ),
                        "type": "website",
                        "chunk": i
                    }
                )

    return documents


# ============================================================
# LOAD KNOWLEDGE BASE
# ============================================================

if not st.session_state.knowledge_loaded:

    with st.spinner(
        "Loading academic documents..."
    ):

        st.session_state.academic_kb = (
            load_academic_documents()
        )

        st.session_state.knowledge_loaded = True


if not st.session_state.website_loaded:

    with st.spinner(
        "Loading official VU information..."
    ):

        st.session_state.website_kb = (
            load_website_documents()
        )

        st.session_state.website_loaded = True


academic_kb = st.session_state.academic_kb
website_kb = st.session_state.website_kb


# ============================================================
# TOKENIZER & RETRIEVAL
# ============================================================

STOPWORDS = {
    "the",
    "a",
    "an",
    "is",
    "are",
    "am",
    "i",
    "me",
    "my",
    "to",
    "of",
    "in",
    "on",
    "for",
    "and",
    "or",
    "can",
    "could",
    "would",
    "should",
    "do",
    "does",
    "did",
    "be",
    "it",
    "this",
    "that",
    "with",
    "from",
    "at",
    "as",
    "what",
    "which",
    "how",
    "where",
    "when",
    "why",
    "you",
    "your"
}


def tokenize(text):

    tokens = re.findall(
        r"[a-zA-Z0-9]+",
        text.lower()
    )

    return [
        token
        for token in tokens
        if token not in STOPWORDS
    ]


def retrieve(
    query,
    documents,
    top_k=6,
    minimum_score=0.03
):

    if not documents:
        return []

    query_tokens = tokenize(
        query
    )

    if not query_tokens:
        return []

    query_counter = Counter(
        query_tokens
    )

    doc_freq = Counter()

    tokenized_docs = []

    for doc in documents:

        tokens = tokenize(
            doc["text"]
        )

        for token in set(tokens):

            doc_freq[token] += 1

        tokenized_docs.append(
            tokens
        )

    total = len(
        documents
    )

    scored = []

    boost_phrases = [
        "minimum cgpa",
        "attendance",
        "eligibility",
        "eligible",
        "prerequisite",
        "minor",
        "semester",
        "admission",
        "course",
        "credits",
        "programme",
        "program",
        "transfer",
        "change degree",
        "change programme",
        "law",
        "summer",
        "summer term",
        "registration"
    ]

    query_lower = query.lower()

    for i, doc in enumerate(
        documents
    ):

        tokens = tokenized_docs[i]

        if not tokens:
            continue

        token_counter = Counter(
            tokens
        )

        score = sum(
            (
                token_counter[token]
                / len(tokens)
            )
            *
            (
                math.log(
                    (total + 1)
                    /
                    (
                        doc_freq.get(
                            token,
                            0
                        ) + 1
                    )
                )
                + 1
            )
            *
            query_count

            for token, query_count
            in query_counter.items()

            if token in token_counter
        )

        text_lower = doc["text"].lower()

        for phrase in boost_phrases:

            if (
                phrase in query_lower
                and phrase in text_lower
            ):

                score += 0.15

        if score >= minimum_score:

            scored.append(
                (
                    score,
                    doc
                )
            )

    scored.sort(
        key=lambda x: x[0],
        reverse=True
    )

    return [
        doc
        for score, doc
        in scored[:top_k]
    ]


# ============================================================
# QUERY CLASSIFICATION
# ============================================================

def classify_query(question):

    q = question.lower()

    website_keywords = [
        "admission",
        "apply",
        "application",
        "contact",
        "phone number",
        "email",
        "campus",
        "location",
        "visit university",
        "website",
        "school of",
        "programme page"
    ]

    if any(
        keyword in q
        for keyword in website_keywords
    ):
        return "website"

    return "academic"


# ============================================================
# EFFECTIVE PROFILE FOR QUESTION
# ============================================================

def effective_profile_for_question(
    question,
    base_profile
):

    profile = dict(
        base_profile
    )

    q = question.lower()

    # Program aliases

    program_aliases = {
        "b.tech": "B.Tech",
        "btech": "B.Tech",
        "b tech": "B.Tech",
        "bms": "BMS",
        "ba llb": "BA LLB",
        "ba, llb": "BA LLB",
        "bms llb": "BMS LLB"
    }

    for alias, value in program_aliases.items():

        if alias in q:

            profile["program"] = value

            break

    # CGPA

    cgpa_match = re.search(
        r"\bcgpa\s*(?:is|of|:)?\s*([0-9]+(?:\.[0-9]+)?)",
        q
    )

    if cgpa_match:

        try:

            profile["cgpa"] = float(
                cgpa_match.group(1)
            )

        except ValueError:
            pass

    # Credits

    credits_match = re.search(
        r"(\d+)\s*(?:credits?|cr)\b",
        q
    )

    if credits_match:

        try:

            profile["completed_credits"] = int(
                credits_match.group(1)
            )

        except ValueError:
            pass

    # Semester

    semester_match = re.search(
        r"\b(\d+)(?:st|nd|rd|th)?\s*semester\b",
        q
    )

    if semester_match:

        profile["semester"] = (
            semester_match.group(1)
            + "th Semester"
        )

    return profile


# ============================================================
# SYSTEM PROMPT
# ============================================================

def build_system_prompt(
    profile,
    retrieved_context
):

    profile_text = profile_to_text(
        profile
    )

    return f"""
You are the AI Academic Advisor for Vidyashilp University.

Your task is to provide accurate, grounded academic guidance.

IMPORTANT SOURCE RULES:

1. Official Vidyashilp University academic documents and
   retrieved official website information are the authority
   for university rules.

2. The synthetic student database is NOT an authority for
   university rules.

3. Use the student profile only to personalize the answer.

4. Never invent prerequisites, credits, CGPA requirements,
   attendance rules, fees, deadlines, policies or course rules.

5. If the retrieved documents do not contain enough information,
   clearly say that the information could not be verified from
   the available university sources.

6. If the student profile conflicts with official university
   documents, explicitly identify the conflict and prioritize
   the official university source.

7. Do not assume that a student is eligible simply because a
   synthetic profile says they are eligible.

8. Verify prerequisites, credits and progression requirements
   against retrieved university documents whenever possible.

9. If information is missing from the profile, do not guess it.

10. If the student has changed degree/programme/course history,
    treat the current programme as the current academic status,
    while using historical information only as context.
    Verify transfer/change-of-programme rules from university
    documents before giving a conclusion.

11. Never reveal information belonging to another student.

12. The student database contains synthetic test data only.

13. For privacy/security questions, do not expose unnecessary
    personal information.

CURRENT STUDENT PROFILE:

{profile_text}

RETRIEVED UNIVERSITY CONTEXT:

{retrieved_context}

ANSWERING STYLE:

- Be clear and concise.
- Give the direct answer first.
- Explain the reasoning when eligibility or planning is involved.
- Mention the relevant source when possible.
- If the answer is uncertain, say so.
- Do not fabricate information.
"""


# ============================================================
# CREATE CONTEXT
# ============================================================

def create_context(
    documents
):

    if not documents:

        return "No relevant university information was retrieved."

    parts = []

    for index, doc in enumerate(
        documents,
        start=1
    ):

        source = doc.get(
            "source",
            "Unknown source"
        )

        chunk = doc.get(
            "chunk",
            ""
        )

        page = doc.get(
            "page"
        )

        section = doc.get(
            "section",
            ""
        )

        source_label = (
            f"[{source}#{chunk}]"
        )

        if page:

            source_label += (
                f" page {page}"
            )

        if section:

            source_label += (
                f" section: {section}"
            )

        parts.append(
            f"{source_label}\n{doc['text']}"
        )

    return "\n\n".join(
        parts
    )


# ============================================================
# CONVERSATION CONTEXT
# ============================================================

def get_recent_conversation():

    if not st.session_state.messages:
        return ""

    recent = st.session_state.messages[-6:]

    lines = []

    for message in recent:

        role = message.get(
            "role",
            ""
        )

        content = message.get(
            "content",
            ""
        )

        if role and content:

            lines.append(
                f"{role.upper()}: {content}"
            )

    return "\n".join(
        lines
    )


# ============================================================
# GENERATE ANSWER
# ============================================================

def generate_answer(
    question,
    profile,
    context
):

    if not api_key:

        return (
            "The Groq API key is not configured. "
            "Please add GROQ_API_KEY to Streamlit secrets."
        )

    client = Groq(
        api_key=api_key
    )

    system_prompt = build_system_prompt(
        profile,
        context
    )

    conversation = get_recent_conversation()

    user_prompt = f"""
Previous conversation:

{conversation}

Current student question:

{question}

Use the retrieved university context and student profile
to answer the current question.
"""

    models = [
        MODEL,
        "openai/gpt-oss-20b",
        "openai/gpt-oss-120b"
    ]

    last_error = None

    for model_name in models:

        try:

            response = client.chat.completions.create(
                model=model_name,
                messages=[
                    {
                        "role": "system",
                        "content": system_prompt
                    },
                    {
                        "role": "user",
                        "content": user_prompt
                    }
                ],
                temperature=0.1,
                max_tokens=900
            )

            answer = response.choices[0].message.content

            if answer:

                return answer.strip()

        except Exception as error:

            last_error = error

    return (
        "I’m unable to generate the answer right now. "
        "Please try again."
    )


# ============================================================
# DISPLAY SOURCES
# ============================================================

def display_sources(
    documents
):

    if not documents:
        return

    with st.expander(
        "📚 Sources used"
    ):

        for doc in documents:

            source = doc.get(
                "source",
                "Unknown"
            )

            chunk = doc.get(
                "chunk",
                ""
            )

            page = doc.get(
                "page"
            )

            section = doc.get(
                "section",
                ""
            )

            details = (
                f"{source}"
                f" — chunk {chunk}"
            )

            if page:
                details += (
                    f" — page {page}"
                )

            if section:
                details += (
                    f" — {section}"
                )

            st.markdown(
                f"- **{details}**"
            )

            if doc.get("url"):

                st.markdown(
                    f"[Open source]({doc['url']})"
                )


# ============================================================
# SUGGESTIONS
# ============================================================

suggestions = [
    "What are the attendance requirements?",
    "How can I check my course eligibility?",
    "What happens if I fail a prerequisite?",
    "How many credits do I need?",
    "Can I change my programme?",
    "How should I plan my next semester?"
]


# ============================================================
# WELCOME
# ============================================================

if not st.session_state.messages:

    st.markdown(
        """
        <div class="info-box">

        <h3>🎓 Welcome to the VU AI Academic Advisor</h3>

        <p>
        Ask questions about courses, credits, prerequisites,
        attendance, registration, programme progression and
        academic planning.
        </p>

        <p>
        For A2 testing, you can select a synthetic student
        profile from the sidebar.
        </p>

        </div>
        """,
        unsafe_allow_html=True
    )

    cols = st.columns(3)

    for i, suggestion in enumerate(
        suggestions[:3]
    ):

        with cols[i]:

            if st.button(
                suggestion,
                use_container_width=True
            ):

                st.session_state.pending_question = (
                    suggestion
                )

                st.rerun()


# ============================================================
# DISPLAY CHAT HISTORY
# ============================================================

for message in st.session_state.messages:

    role = message.get(
        "role"
    )

    content = message.get(
        "content",
        ""
    )

    if role == "user":

        st.markdown(
            f"""
            <div class="bubble-user">
                <div class="bubble-user-inner">
                    {html.escape(content)}
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

    elif role == "assistant":

        st.markdown(
            f"""
            <div class="bubble-bot">
                <div class="bubble-bot-inner">
                    {content}
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )


# ============================================================
# CHAT INPUT
# ============================================================

question = st.chat_input(
    "Ask your academic question..."
)

if (
    not question
    and st.session_state.pending_question
):

    question = (
        st.session_state.pending_question
    )

    st.session_state.pending_question = None


# ============================================================
# PROCESS QUESTION
# ============================================================

if question:

    question = question.strip()

    if not question:
        st.stop()

    st.session_state.messages.append(
        {
            "role": "user",
            "content": question
        }
    )

    # --------------------------------------------------------
    # Greeting / social
    # --------------------------------------------------------

    if SOCIAL_RE.match(question):

        if re.search(
            r"\b(thank|thanks|thx|ty)\b",
            question,
            re.I
        ):

            answer = THANKS_REPLY

        else:

            answer = GREETING_REPLY

        st.session_state.messages.append(
            {
                "role": "assistant",
                "content": answer
            }
        )

        st.rerun()

    # --------------------------------------------------------
    # Personal
    # --------------------------------------------------------

    if PERSONAL_RE.search(
        question
    ):

        answer = PERSONAL_REPLY

        st.session_state.messages.append(
            {
                "role": "assistant",
                "content": answer
            }
        )

        st.rerun()

    # --------------------------------------------------------
    # Out of scope
    # --------------------------------------------------------

    if OUT_OF_SCOPE_RE.search(
        question
    ):

        answer = OUT_OF_SCOPE_REPLY

        st.session_state.messages.append(
            {
                "role": "assistant",
                "content": answer
            }
        )

        st.rerun()

    # --------------------------------------------------------
    # Profile
    # --------------------------------------------------------

    effective_profile = (
        effective_profile_for_question(
            question,
            student_profile
        )
    )

    # --------------------------------------------------------
    # Query category
    # --------------------------------------------------------

    category = classify_query(
        question
    )

    # --------------------------------------------------------
    # Retrieval
    # --------------------------------------------------------

    retrieval_query = question

    profile_terms = (
        profile_retrieval_terms(
            effective_profile
        )
    )

    profile_sensitive_words = [
        "eligible",
        "eligibility",
        "prerequisite",
        "course",
        "courses",
        "credits",
        "minor",
        "registration",
        "semester",
        "plan",
        "planning",
        "failed",
        "change",
        "transfer"
    ]

    if (
        category == "academic"
        and any(
            word in question.lower()
            for word in profile_sensitive_words
        )
        and profile_terms
    ):

        retrieval_query = (
            question
            + " "
            + profile_terms
        )

    if category == "website":

        retrieved_documents = retrieve(
            retrieval_query,
            website_kb,
            top_k=6
        )

    else:

        retrieved_documents = retrieve(
            retrieval_query,
            academic_kb,
            top_k=6
        )

    # --------------------------------------------------------
    # If no source
    # --------------------------------------------------------

    if not retrieved_documents:

        answer = (
            "I couldn't find enough relevant information "
            "in the available Vidyashilp University sources "
            "to answer this reliably. I don't want to guess."
        )

        st.session_state.messages.append(
            {
                "role": "assistant",
                "content": answer
            }
        )

        st.rerun()

    # --------------------------------------------------------
    # Context
    # --------------------------------------------------------

    context = create_context(
        retrieved_documents
    )

    # --------------------------------------------------------
    # Generate answer
    # --------------------------------------------------------

    with st.spinner(
        "Checking the university information..."
    ):

        answer = generate_answer(
            question,
            effective_profile,
            context
        )

    # --------------------------------------------------------
    # Save assistant response
    # --------------------------------------------------------

    st.session_state.messages.append(
        {
            "role": "assistant",
            "content": answer
        }
    )

    # --------------------------------------------------------
    # Display answer
    # --------------------------------------------------------

    st.markdown(
        f"""
        <div class="bubble-bot">
            <div class="bubble-bot-inner">
                {answer}
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

    display_sources(
        retrieved_documents
    )

    st.rerun()


# ============================================================
# FOOTER
# ============================================================

st.markdown(
    """
    <div class="footer-text">
        Vidyashilp University · AI Academic Advisor ·
        DATA308 Generative AI
    </div>
    """,
    unsafe_allow_html=True
)
